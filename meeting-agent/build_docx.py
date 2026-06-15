"""Генерация .docx-протокола из структурированного JSON.

Возвращает файл в памяти (BytesIO) для st.download_button и имя файла.
Опционально сохраняет копию в output/.
"""

from __future__ import annotations

import datetime as _dt
import io
import os
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def _today() -> str:
    return _dt.date.today().isoformat()


def _safe_date_for_filename(date_value: str | None) -> str:
    """Подобрать дату для имени файла: ISO-дата из JSON или сегодняшняя."""
    if date_value and re.search(r"\d{4}-\d{2}-\d{2}", str(date_value)):
        return re.search(r"\d{4}-\d{2}-\d{2}", str(date_value)).group(0)
    return _today()


def _add_heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def _add_bullets(doc: Document, items: list[str]) -> None:
    if not items:
        doc.add_paragraph("— нет —")
        return
    for item in items:
        text = str(item).strip()
        if text:
            doc.add_paragraph(text, style="List Bullet")


def build_docx(data: dict[str, Any]) -> tuple[io.BytesIO, str]:
    """Собрать .docx из протокола. Вернуть (BytesIO, имя_файла)."""
    doc = Document()

    # Шапка.
    title = doc.add_heading("ПРОТОКОЛ СОВЕЩАНИЯ", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    тема = str(data.get("тема") or "").strip()
    дата = str(data.get("дата") or "").strip() or "не указана"
    участники = data.get("участники") or []

    if тема:
        p = doc.add_paragraph()
        run = p.add_run(f"Тема: {тема}")
        run.bold = True
        run.font.size = Pt(12)

    doc.add_paragraph(f"Дата: {дата}")
    if участники:
        doc.add_paragraph("Участники: " + ", ".join(str(u).strip() for u in участники))

    # 1. Повестка.
    _add_heading(doc, "1. Повестка", level=1)
    _add_bullets(doc, data.get("повестка") or [])

    # 2. Решения.
    _add_heading(doc, "2. Решения", level=1)
    _add_bullets(doc, data.get("решения") or [])

    # 3. Поручения (таблица).
    _add_heading(doc, "3. Поручения", level=1)
    поручения = data.get("поручения") or []
    if not поручения:
        doc.add_paragraph("— нет —")
    else:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        headers = ("Задача", "Ответственный", "Срок", "Приоритет")
        for cell, header in zip(table.rows[0].cells, headers):
            cell.paragraphs[0].add_run(header).bold = True
        for task in поручения:
            row = table.add_row().cells
            row[0].text = str(task.get("задача") or "").strip()
            row[1].text = str(task.get("ответственный") or "не определён").strip()
            row[2].text = str(task.get("срок") or "не указан").strip()
            row[3].text = str(task.get("приоритет") or "средний").strip()

    # 4. Открытые вопросы.
    _add_heading(doc, "4. Открытые вопросы", level=1)
    _add_bullets(doc, data.get("открытые_вопросы") or [])

    # 5. Следующая встреча.
    _add_heading(doc, "5. Следующая встреча", level=1)
    следующая = str(data.get("следующая_встреча") or "").strip()
    doc.add_paragraph(следующая if следующая else "— не назначена —")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    filename = f"Протокол_{_safe_date_for_filename(data.get('дата'))}.docx"
    return buffer, filename


def save_to_output(buffer: io.BytesIO, filename: str, output_dir: str = "output") -> str:
    """Сохранить копию .docx в output/. Вернуть путь к файлу."""
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, filename)
    with open(path, "wb") as f:
        f.write(buffer.getvalue())
    buffer.seek(0)
    return path
